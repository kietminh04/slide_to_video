import io
import fitz  # PyMuPDF
from pptx import Presentation
from docx import Document

class ImageSlideExtractor:
    @staticmethod
    def extract_from_pdf(file_bytes: bytes) -> list[bytes]:
        """Render từng trang PDF thành ảnh PNG 150 DPI hoàn toàn trên RAM."""
        images = []
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            pix = page.get_pixmap(dpi=150)
            images.append(pix.tobytes("png"))
        return images

    @staticmethod
    def extract_from_pptx(file_bytes: bytes) -> list[bytes]:
        """Trích xuất ảnh nhúng toàn trang từ các slide PPTX."""
        images = []
        prs = Presentation(io.BytesIO(file_bytes))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == 13:  # Picture
                    image = shape.image
                    images.append(image.blob)
                    break  # Lấy ảnh đại diện chính của slide
        return images

    @staticmethod
    def extract_from_docx(file_bytes: bytes) -> list[bytes]:
        """Trích xuất ảnh nhúng từ DOCX."""
        images = []
        doc = Document(io.BytesIO(file_bytes))
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                images.append(rel.target_part.blob)
        return images
