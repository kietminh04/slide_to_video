"""Engine điều phối Multi-Agent Pipeline tự đánh giá & xuất kịch bản đa định dạng.

Quy trình khép kín:
1. Extractor: Bóc tách đa định dạng (PPTX/PDF/DOCX/TXT).
2. Planner: Tính toán ngân sách từ toán học & phân bổ Bloom.
3. Generator: Sinh kịch bản Scene và Beat (Rolling Context).
4. Critic Agent: Tự đánh giá 4 trục (Pacing, Terminology, Grounding, Visual).
5. Self-Correction Loop: Refiner Agent sửa lỗi có chọn lọc (Tối đa 2 vòng lặp).
6. Hoàn thiện Timeline & Xuất đa định dạng (JSON, Markdown, Word, SRT).
"""

from __future__ import annotations

import json
from pathlib import Path

from codebase.clsg.config import Settings, get_settings
from codebase.clsg.critic import CriticAgent
from codebase.clsg.extract import extract
from codebase.clsg.generate import generate as generate_scenes
from codebase.clsg.glossary import PedagogicalGlossary
from codebase.clsg.llm import FakeLLMProvider, LLMProvider, OpenAIProvider
from codebase.clsg.plan import plan
from codebase.clsg.refiner import RefinerAgent
from codebase.clsg.schemas import (
    Chunk,
    CriticScorecard,
    Script,
    ScriptMeta,
    WeightProfile,
)
from codebase.clsg.style import apply_style


def _compute_timeline(script: Script, wpm: int = 140, target_duration_s: float | None = None) -> Script:
    """Tính thời lượng chuẩn hóa sử dụng ProsodyPlanner và ActiveRescaler."""
    from codebase.clsg.guard import ActiveRescaler
    from codebase.clsg.prosody import ProsodyPlanner

    if target_duration_s and target_duration_s > 0:
        script, _ = ActiveRescaler.auto_rescale(script, target_duration_s=target_duration_s, base_wpm=wpm)
        return script

    planner = ProsodyPlanner(wpm=wpm)
    t = 0.0

    for scene in script.scenes:
        scene.t_start = round(t, 2)
        scene_duration = 0.0
        for beat in scene.beats:
            res = planner.plan(beat.display_text, custom_wpm=wpm)
            beat.est_s = res.calibrated_duration_s
            beat.tts_text = res.marked_text
            scene_duration += res.calibrated_duration_s
        scene.t_end = round(t + scene_duration, 2)
        t = scene.t_end

    return script


def run_pipeline(
    doc_path: Path,
    weights_path: Path | None = None,
    *,
    fake: bool = False,
    style_pack: str = "serious",
    settings: Settings | None = None,
    max_refine_rounds: int = 2,
) -> Script:
    """Chạy toàn bộ pipeline Multi-Agent tự đánh giá.

    Args:
        doc_path: Đường dẫn tài liệu (PPTX, DOCX, PDF, TXT).
        weights_path: Đường dẫn file trọng số JSON (tùy chọn).
        fake: Nếu True, dùng FakeLLMProvider (cho test).
        style_pack: Tên style pack (serious, drama_tongtai...).
        settings: Cấu hình pipeline.
        max_refine_rounds: Số vòng lặp tự sửa lỗi tối đa (mặc định 2).

    Returns:
        Script hoàn chỉnh kèm bảng điểm CriticScorecard.
    """
    settings = settings or get_settings()
    glossary = PedagogicalGlossary()

    # === 1. Khởi tạo LLM Provider ===
    if fake:
        llm: LLMProvider = FakeLLMProvider()
    else:
        llm = OpenAIProvider(
            api_key=settings.openai_api_key,
            base_url=settings.openai_base_url,
            cache_dir=settings.llm_cache_dir,
            usage_log=settings.usage_log_file,
            default_model=settings.model_generator,
        )

    # === 2. Bóc tách tài liệu (Extractor Agent) ===
    chunks: list[Chunk] = extract(doc_path)
    if not chunks:
        raise ValueError(f"Không trích xuất được nội dung nào từ {doc_path}")

    # === 3. Đọc hoặc tự sinh Trọng số (Planner Agent) ===
    all_chapters = sorted(list({c.chapter for c in chunks}))
    if weights_path and weights_path.exists():
        weights_data = json.loads(weights_path.read_text(encoding="utf-8"))
        weights = WeightProfile(**weights_data)
    else:
        # Tự động sinh trọng số đồng đều nếu không có file weights
        equal_w = 1.0 / len(all_chapters)
        weights = WeightProfile(
            chapters=all_chapters,
            w_user=[equal_w] * len(all_chapters)
        )

    # === 4. Lập kế hoạch sư phạm (Pedagogical Planner) ===
    blueprint = plan(chunks, weights, settings, llm)

    # === 5. Sinh kịch bản đa luồng (Multi-Track Generator) ===
    scenes, fact_sheets = generate_scenes(blueprint, chunks, llm, settings)

    weight_dict = {seg.chapter: seg.weight for seg in blueprint.segments}
    script = Script(
        meta=ScriptMeta(
            doc_id=doc_path.stem,
            duration_s=int(settings.minutes * 60),
            style=style_pack,
            weights=weight_dict,
        ),
        scenes=scenes,
    )

    # === 6. Vòng Lặp Phản Biện & Tự Sửa Lỗi (Critic-Refiner Loop) ===
    critic = CriticAgent(glossary=glossary, settings=settings)
    refiner = RefinerAgent(glossary=glossary)

    current_scorecard = critic.evaluate(script, blueprint, chunks)
    rounds = 0

    while not current_scorecard.passed and rounds < max_refine_rounds:
        rounds += 1
        script = refiner.refine(script, blueprint, chunks, current_scorecard)
        current_scorecard = critic.evaluate(script, blueprint, chunks)

    script.critic_scorecard = current_scorecard

    # === 7. Áp dụng Style Pack & Phân bổ thời gian ===
    script = apply_style(script, style_pack, llm, settings)
    script = _compute_timeline(script, wpm=settings.wpm, target_duration_s=script.meta.duration_s)

    return script


# === Các hàm xuất bản kịch bản đa định dạng (Export Utilities) ===

def export_markdown(script: Script, output_path: Path) -> Path:
    """Xuất kịch bản ra định dạng Markdown Transcript chuyên nghiệp."""
    total_words = sum(len(b.display_text.split()) for s in script.scenes for b in s.beats)
    scorecard = script.critic_scorecard

    lines = [
        f"# KỊCH BẢN BÀI GIẢNG: {script.meta.doc_id.upper()}",
        f"",
        f"> **Thời lượng dự kiến:** {script.meta.duration_s // 60}m {script.meta.duration_s % 60}s ({script.meta.duration_s}s)  ",
        f"> **Tổng số từ:** {total_words} từ ({script.meta.style})  ",
        f"> **Số phân cảnh:** {len(script.scenes)} slides  ",
    ]

    if scorecard:
        lines.append(
            f"> **Điểm đánh giá Critic:** {scorecard.overall_score}/100 "
            f"(Pacing: {scorecard.pacing_score} | Thuật ngữ: {scorecard.terminology_score} | "
            f"Grounding: {scorecard.grounding_score} | Visual: {scorecard.visual_score})  "
        )

    lines.extend([
        "",
        "---",
        "",
        "## BẢNG PHÂN CẢNH CHI TIẾT (TIMELINE TRANSCRIPT)",
        "",
        "| Thời gian | Phân cảnh | Lời giảng sư phạm (Narration) | Chỉ dẫn thị giác (Visual Intent) |",
        "| :--- | :--- | :--- | :--- |",
    ])

    for s in script.scenes:
        time_label = f"{int(s.t_start // 60):02d}:{int(s.t_start % 60):02d} - {int(s.t_end // 60):02d}:{int(s.t_end % 60):02d}"
        dur_label = f"{s.t_end - s.t_start:.1f}s"
        narration = " ".join(b.display_text for b in s.beats).replace("|", "-")
        v_type = s.visual.visual_type or s.visual.layout or "Mặc định"
        v_purpose = s.visual.visual_purpose or s.visual.asset_hint or ""
        visual_desc = f"**[{v_type}]**<br/>{v_purpose}" if v_purpose else f"**[{v_type}]**"

        lines.append(f"| `{time_label}` ({dur_label}) | **{s.id}** ({s.depth}) | {narration} | {visual_desc} |")

    lines.extend([
        "",
        "---",
        "*Tài liệu được sinh tự động bởi Hệ thống CLSG Multi-Agent.*"
    ])

    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def export_word_doc(script: Script, output_path: Path) -> Path:
    """Xuất kịch bản ra định dạng Word Document (.docx) bảng phân cảnh 3 cột."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # Tiêu đề bài giảng
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f"KỊCH BẢN PHÂN CẢNH BÀI GIẢNG: {script.meta.doc_id.upper()}")
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    # Thông số metadata
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(
        f"Thời lượng: {script.meta.duration_s}s | Số slide: {len(script.scenes)} | Phong cách: {script.meta.style}"
    )

    # Bảng phân cảnh 3 cột
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Thời gian & Cảnh"
    hdr_cells[1].text = "Lời giảng Sư phạm (Audio Narration)"
    hdr_cells[2].text = "Chỉ dẫn Thị giác (Visual Intent)"

    for s in script.scenes:
        row_cells = table.add_row().cells
        time_label = f"{int(s.t_start // 60):02d}:{int(s.t_start % 60):02d} - {int(s.t_end // 60):02d}:{int(s.t_end % 60):02d}"
        row_cells[0].text = f"{time_label}\n({s.id})\n[{s.depth}]"
        row_cells[1].text = " ".join(b.display_text for b in s.beats)
        v_type = s.visual.visual_type or "title+bullets"
        v_purpose = s.visual.visual_purpose or ""
        row_cells[2].text = f"[{v_type}]\n{v_purpose}"

    doc.save(str(output_path))
    return output_path


def export_srt(script: Script, output_path: Path) -> Path:
    """Xuất phụ đề định dạng chuẩn .srt."""
    def fmt_srt(sec: float) -> str:
        h = int(sec // 3600)
        m = int((sec % 3600) // 60)
        s = int(sec % 60)
        ms = int((sec % 1) * 1000)
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"

    lines = []
    idx = 1
    for scene in script.scenes:
        t_cur = scene.t_start
        for beat in scene.beats:
            dur = beat.est_s or 5.0
            t_next = t_cur + dur
            lines.append(f"{idx}")
            lines.append(f"{fmt_srt(t_cur)} --> {fmt_srt(t_next)}")
            lines.append(beat.display_text)
            lines.append("")
            idx += 1
            t_cur = t_next

    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path
